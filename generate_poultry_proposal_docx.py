"""Generates a properly formatted Word document for the poultry youth-loan project proposal."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x2F, 0x52, 0x33)       # deep veld green
ACCENT_SOFT = "E7EDE3"                     # light green tint (hex, no #)
INK = RGBColor(0x20, 0x1E, 0x19)
INK_SOFT = RGBColor(0x5A, 0x56, 0x4C)
INK_FAINT = RGBColor(0x8C, 0x85, 0x77)
LINE_GREY = "DDD7C8"

HEADING_FONT = "Cambria"
BODY_FONT = "Calibri"

doc = Document()

# ---- Page setup ----
section = doc.sections[0]
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)

# ---- Base style ----
normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(11)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.25


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, color=LINE_GREY, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def style_table(table, header_row=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_borders(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = BODY_FONT
            if header_row and r_idx == 0:
                set_cell_shading(cell, ACCENT_SOFT)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = ACCENT


def add_heading(text, number):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    run_num = p.add_run(f"{number}  ")
    run_num.font.name = HEADING_FONT
    run_num.font.size = Pt(11)
    run_num.font.color.rgb = ACCENT
    run_num.font.bold = True
    run_txt = p.add_run(text.upper())
    run_txt.font.name = HEADING_FONT
    run_txt.font.size = Pt(14)
    run_txt.font.bold = True
    run_txt.font.color.rgb = INK
    # thin accent rule under heading
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(8)
    rule_run = rule.add_run("―" * 18)
    rule_run.font.color.rgb = ACCENT
    rule_run.font.size = Pt(8)


def add_body(text, italic=False, bold=False, color=None, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(11)
    run.font.italic = italic
    run.font.bold = bold
    run.font.color.rgb = color or INK
    return p


def add_bullets(items, ordered=False):
    for item in items:
        p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        run.font.name = BODY_FONT
        run.font.size = Pt(11)
        run.font.color.rgb = INK


def add_field_row(table, label, value, confirmed=False):
    row = table.add_row()
    row.cells[0].text = ""
    p0 = row.cells[0].paragraphs[0]
    r0 = p0.add_run(label)
    r0.font.bold = True
    r0.font.size = Pt(10)
    r0.font.color.rgb = INK_SOFT
    row.cells[1].text = ""
    p1 = row.cells[1].paragraphs[0]
    r1 = p1.add_run(value)
    r1.font.italic = not confirmed
    r1.font.bold = confirmed
    r1.font.size = Pt(10)
    r1.font.color.rgb = INK if confirmed else INK_FAINT


# ============================================================
# TITLE BLOCK
# ============================================================
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
eyebrow = title_p.add_run("PROJECT PROPOSAL — YOUTH LOAN APPLICATION\n")
eyebrow.font.name = HEADING_FONT
eyebrow.font.italic = True
eyebrow.font.size = Pt(11)
eyebrow.font.color.rgb = ACCENT
title_p.paragraph_format.space_after = Pt(2)

main_title = doc.add_paragraph()
main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
mt_run = main_title.add_run("Broiler Poultry Production Enterprise")
mt_run.font.name = HEADING_FONT
mt_run.font.size = Pt(24)
mt_run.font.bold = True
mt_run.font.color.rgb = ACCENT
main_title.paragraph_format.space_after = Pt(10)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta_text = ("Presented via the office of Hon. Discent Collins Bajila, MP  |  "
             "Venue: Luveve Youth Centre  |  Date: 12 August 2026, 09:30")
meta_run = meta.add_run(meta_text)
meta_run.font.size = Pt(10)
meta_run.font.color.rgb = INK_SOFT
meta.paragraph_format.space_after = Pt(20)

# Bottom border under the title block
border_p = doc.add_paragraph()
border_p.paragraph_format.space_after = Pt(14)
pPr = border_p._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "18")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), "2F5233")
pBdr.append(bottom)
pPr.append(pBdr)

# ============================================================
# APPLICANT DETAILS TABLE
# ============================================================
field_table = doc.add_table(rows=0, cols=2)
field_table.autofit = True
add_field_row(field_table, "Full Name", "Bhekimpilo Nyoni")
add_field_row(field_table, "National ID Number", "ID number")
add_field_row(field_table, "Physical Address / Ward", "6658 Gwabalanda, PO Luveve, Bulawayo")
add_field_row(field_table, "Phone Number", "Phone number")
add_field_row(field_table, "Business Name", "Nyoni Poultry Enterprise")
for row in field_table.rows:
    for cell in row.cells:
        set_cell_borders(cell)
        cell.width = Inches(3.2)

# ============================================================
# 01 EXECUTIVE SUMMARY
# ============================================================
add_heading("Executive Summary", "01")
add_body(
    "This proposal seeks funding to start a small-scale broiler chicken production business, "
    "beginning with a first batch of 100 day-old broiler chicks. Broiler production was chosen "
    "because it has a short production cycle (6\u20138 weeks from chick to sale), strong and "
    "steady local demand for chicken meat, and a proven, well-understood process a first-time "
    "farmer can manage successfully with basic guidance."
)

stat_table = doc.add_table(rows=2, cols=3)
stat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
labels = ["Capital Requested", "Production Cycle", "Projected Revenue"]
values = ["US$464.00", "6\u20138 weeks", "US$630.00"]
for i in range(3):
    c = stat_table.cell(0, i)
    c.text = ""
    r = c.paragraphs[0].add_run(labels[i])
    r.font.size = Pt(9)
    r.font.color.rgb = INK_FAINT
    r.font.bold = True
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_shading(c, ACCENT_SOFT)
    set_cell_borders(c)

    c2 = stat_table.cell(1, i)
    c2.text = ""
    r2 = c2.paragraphs[0].add_run(values[i])
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = ACCENT
    c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_shading(c2, ACCENT_SOFT)
    set_cell_borders(c2)

add_body(
    "Figures shown are the estimate used throughout this proposal \u2014 see Section 06 for the "
    "full itemized budget and the note on replacing estimates with real supplier quotations.",
    italic=True, color=INK_SOFT, space_after=4
)

# ============================================================
# 02 ABOUT THE APPLICANT
# ============================================================
add_heading("About the Applicant", "02")
add_body(
    "I am a young resident of Luveve who has always been drawn to poultry farming as a realistic "
    "way to build my own income rather than wait indefinitely for formal employment. I already "
    "have a small structure at home that can house between 50 and 100 chicks, which I set up "
    "myself with limited resources \u2014 this loan will let me extend that structure and add "
    "separate compartments so I can rotate batches, keeping incoming day-old chicks apart from an "
    "older, near-market-ready batch. That way I can sell one batch and start the next immediately, "
    "without any downtime and without risking young chicks by mixing them with older birds. "
    "Poultry farming appeals to me because the short production cycle means I can see real results "
    "and start repaying quickly, and because it lets me serve my own community directly \u2014 "
    "supplying fresh, affordable chicken to my neighbours here in Luveve. My family also has a "
    "rural home about 60km from Bulawayo, which gives me a natural second site to expand "
    "production and sales to once the Luveve operation is established, and I plan to use solar "
    "power to run the project reliably at both sites without depending on an unstable grid "
    "supply. As the business grows, I also want to create an opportunity to bring other young "
    "people from the area on board."
)

# ============================================================
# 03 BUSINESS DESCRIPTION
# ============================================================
add_heading("Business Description", "03")
add_body("Nature of business: small-scale commercial broiler production for sale of live birds "
          "to local households, vendors, and butcheries.", bold=False)
add_body("Objectives:", bold=True, space_after=2)
add_bullets([
    "Successfully raise and sell the first batch of 100 broilers within 8 weeks of receiving chicks.",
    "Use profit from the first batch to part-repay the loan and fund a second, larger batch (150\u2013200 birds) within 3 months.",
    "Establish a reliable local customer base within the first two production cycles.",
    "Grow into a stable enterprise providing significant income within 12 months.",
    "Expand production to a family rural home approximately 60km from Bulawayo as a second site, "
    "once the Luveve operation is established, opening a second market and spreading production "
    "across two locations.",
], ordered=True)
add_body(
    "Broilers were chosen over layers or crop farming specifically because the short cycle means "
    "faster loan repayment, and chicken meat demand is consistent year-round \u2014 not seasonal.",
    color=INK_SOFT
)

# ============================================================
# 04 MARKET ANALYSIS
# ============================================================
add_heading("Market Analysis", "04")
add_body("Target customers: households in the local area buying live birds directly, local "
          "vendors/tuck-shops, at least one local butchery approached once volumes grow, and a "
          "second customer base at the family's rural home approximately 60km from Bulawayo once "
          "operations expand to that site (Section 03, Objective 5).")
add_body("Demand: chicken is the most affordable, widely consumed meat in Zimbabwe. Local demand "
          "in high-density suburbs consistently outstrips small-scale supply \u2014 birds are "
          "typically pre-ordered by word of mouth before reaching maturity.")
add_body("Pricing: live broilers currently sell locally at approximately US$6\u20138 per bird. "
          "This proposal uses a conservative US$7 per bird throughout.")

# ============================================================
# 05 OPERATIONS PLAN
# ============================================================
add_heading("Operations Plan", "05")
ops_table = doc.add_table(rows=1, cols=2)
ops_table.rows[0].cells[0].text = "Week"
ops_table.rows[0].cells[1].text = "Activity"
ops_rows = [
    ("0", "Receive 100 day-old chicks; brooding house pre-warmed and ready"),
    ("1\u20132", "Brooding \u2014 heat lamp maintained, starter mash, first vaccinations (Newcastle disease, Gumboro)"),
    ("3\u20135", "Grower stage \u2014 grower mash, reduced heating, regular health checks"),
    ("6\u20138", "Finisher stage \u2014 finisher mash, birds reach market weight (1.8\u20132.2kg)"),
    ("8", "Sale of birds; house cleaned, rested, and disinfected before next batch"),
]
for wk, activity in ops_rows:
    row = ops_table.add_row()
    row.cells[0].text = wk
    row.cells[1].text = activity
style_table(ops_table)
ops_table.columns[0].width = Inches(0.8)

add_body(
    "Housing: an existing small structure at the applicant's home, currently sized for 50\u2013100 "
    "chicks, will be extended and divided into two compartments \u2014 one for an incoming batch "
    "of day-old chicks, one for the current batch as it grows toward market weight. This rotation "
    "design means a new batch can start the moment the previous one is sold, without a gap in "
    "production and without mixing different-age birds (which raises disease risk).",
    space_after=6
)
add_body(
    "Power supply: solar power will be used to run the brooding heat lamp and lighting, keeping "
    "the project running reliably during grid outages/load-shedding — brooding chicks are "
    "especially vulnerable to losing heat overnight, so this removes a real risk to the whole "
    "batch. The same solar setup is intended to support the rural home site once operations "
    "expand there.",
    space_after=6
)
add_body("Biosecurity: vaccination schedule followed per supplier/extension officer guidance; "
          "footbath at house entrance; visitors limited during brooding; sick birds isolated "
          "immediately.", color=INK_SOFT)

# ============================================================
# 06 STARTUP BUDGET
# ============================================================
add_heading("Startup Budget \u2014 First Batch", "06")

callout = doc.add_paragraph()
callout.paragraph_format.space_after = Pt(10)
c_run = callout.add_run(
    "Before the meeting: figures below are realistic local estimates. Bring actual written "
    "quotations from a stockfeeds supplier, hatchery, and hardware store (item 6 on the loan "
    "requirements) and update this table with the real numbers."
)
c_run.font.italic = True
c_run.font.size = Pt(10)
c_run.font.color.rgb = ACCENT

budget_rows = [
    ("Day-old broiler chicks", "100", "$0.60", "$60.00"),
    ("Starter mash (50kg)", "2 bags", "$18.00", "$36.00"),
    ("Grower mash (50kg)", "3 bags", "$17.00", "$51.00"),
    ("Finisher mash (50kg)", "3 bags", "$16.00", "$48.00"),
    ("Vaccines (Newcastle + Gumboro)", "1 course", "$15.00", "$15.00"),
    ("Vitamins / electrolytes", "1 course", "$10.00", "$10.00"),
    ("Feeders", "5", "$8.00", "$40.00"),
    ("Drinkers", "5", "$6.00", "$30.00"),
    ("Heat lamp + fitting", "2", "$12.00", "$24.00"),
    ("Wood-shaving litter", "4 bags", "$5.00", "$20.00"),
    ("Structure extension + compartment dividers (wire mesh, poles, roofing offcuts)", "1 set", "$110.00", "$110.00"),
    ("Transport", "\u2014", "\u2014", "$20.00"),
]
budget_table = doc.add_table(rows=1, cols=4)
hdr = budget_table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Item", "Qty", "Unit Cost", "Total"
for item, qty, unit, total in budget_rows:
    row = budget_table.add_row()
    row.cells[0].text = item
    row.cells[1].text = qty
    row.cells[2].text = unit
    row.cells[3].text = total
    for i in (1, 2, 3):
        row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

total_row = budget_table.add_row()
total_row.cells[0].text = "TOTAL ESTIMATED CAPITAL REQUIRED"
total_row.cells[1].text = ""
total_row.cells[2].text = ""
total_row.cells[3].text = "$464.00"
total_row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
style_table(budget_table)
for cell in total_row.cells:
    set_cell_shading(cell, ACCENT_SOFT)
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = ACCENT

add_body(
    "The existing structure already covers the base shelter \u2014 this budget only funds "
    "extending it and adding the two-compartment divider described in Section 05, which is why "
    "the housing line is lower than a full build from scratch would be.",
    italic=True, color=INK_SOFT, space_after=4
)

# ============================================================
# 07 REVENUE PROJECTION
# ============================================================
add_heading("Revenue Projection \u2014 First Batch", "07")
rev_rows = [
    ("Chicks purchased", "100"),
    ("Expected survival rate", "90%"),
    ("Birds available for sale", "90"),
    ("Average selling price / bird", "$7.00"),
    ("Total projected revenue", "$630.00"),
    ("Less: total production cost", "$464.00"),
]
rev_table = doc.add_table(rows=0, cols=2)
for label, value in rev_rows:
    row = rev_table.add_row()
    row.cells[0].text = label
    row.cells[1].text = value
    row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
profit_row = rev_table.add_row()
profit_row.cells[0].text = "Projected net profit, first batch"
profit_row.cells[1].text = "\u2248 $166.00"
profit_row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
style_table(rev_table, header_row=False)
for cell in profit_row.cells:
    set_cell_shading(cell, ACCENT_SOFT)
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = ACCENT

add_body(
    "Profit is modest on the first batch because it also covers the one-off structure extension "
    "and equipment costs reused for every future batch at no extra cost. From the second batch "
    "onward, with the compartments and equipment already in place, the same 100-bird cycle costs "
    "roughly $260 (chicks, feed, vaccines, vitamins, litter, transport only) against the same "
    "~$630 revenue \u2014 a projected profit of \u2248 $370 per cycle. Because the compartment design lets "
    "a fresh batch start the moment the current one is sold, this profit is realistically earned "
    "roughly every 8 weeks on an ongoing basis, not just once.",
    color=INK_SOFT
)

# ============================================================
# 08 LOAN REPAYMENT PLAN
# ============================================================
add_heading("Loan Repayment Plan", "08")
add_bullets([
    "Loan amount requested: US$464.00 (see Section 06), adjusted once real supplier quotations are confirmed",
    "Source of repayment: proceeds from the sale of the first batch of broilers",
    "Proposed terms: insert once known from loan officers at the meeting",
    "Commitment: proceeds from the first batch prioritize loan repayment before any reinvestment into a second batch",
])

# ============================================================
# 09 RISKS & MITIGATION
# ============================================================
add_heading("Risks & Mitigation", "09")
risk_table = doc.add_table(rows=1, cols=2)
risk_table.rows[0].cells[0].text = "Risk"
risk_table.rows[0].cells[1].text = "Mitigation"
risk_rows = [
    ("Disease outbreak", "Strict vaccination schedule, biosecurity measures, guidance from local extension officer"),
    ("High mortality in brooding", "Reliable heat source, reputable hatchery, close monitoring in first two weeks"),
    ("Feed price increases", "Compare prices across at least two local suppliers before purchase"),
    ("Slow sales", "Pre-orders from household contacts before birds reach maturity"),
    ("Theft / predators", "Secure fencing, lockable structure, checked daily"),
    ("Water shortage", "Confirm applicant's water access and backup plan"),
    ("Power outages / load-shedding", "Solar power used for the brooding heat lamp and lighting, reducing dependence on an unstable grid supply"),
]
for risk, mitigation in risk_rows:
    row = risk_table.add_row()
    row.cells[0].text = risk
    row.cells[1].text = mitigation
style_table(risk_table)

# ============================================================
# 10 CONCLUSION
# ============================================================
add_heading("Conclusion", "10")
add_body(
    "This proposal outlines a realistic, low-risk first step into commercial poultry farming, "
    "matched to the scale of support available through the Youth Loan scheme. The short "
    "production cycle means the loan can be repaid quickly from real sales, and the equipment and "
    "skills built up from this first batch create a foundation for a larger, self-sustaining "
    "business \u2014 turning a modest first loan into a real source of income."
)

# ============================================================
# SIGNATURE BLOCK
# ============================================================
doc.add_paragraph().paragraph_format.space_after = Pt(20)
sig_table = doc.add_table(rows=2, cols=2)
sig_table.rows[0].cells[0].text = "_" * 30
sig_table.rows[0].cells[1].text = "_" * 30
sig_table.rows[1].cells[0].text = "Applicant signature & date"
sig_table.rows[1].cells[1].text = "Witness / MP office representative"
for row in sig_table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = INK_SOFT

out_path = r"C:\POC\Personal_Processes\Broiler-Poultry-Project-Proposal.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
